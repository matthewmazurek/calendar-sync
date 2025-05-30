import io
import os
import shutil
import tempfile
from pathlib import Path
from unittest.mock import patch

import pytest
from docx.api import Document

from app import create_app


@pytest.fixture
def app():
    return create_app()


@pytest.fixture
def client(app):
    return app.test_client()


def create_test_docx():
    """Create a simple test .docx file with a basic calendar."""
    doc = Document()
    table = doc.add_table(rows=8, cols=1)
    # Add a month header
    table.cell(0, 0).text = "JANUARY 2025"
    # Add a day with an event
    table.cell(1, 0).text = "1\nNew Year's Day"
    return doc


def save_to_temp_file(doc, extension):
    """Save a document to a temporary file with the given extension."""
    with tempfile.NamedTemporaryFile(suffix=extension, delete=False) as temp:
        doc.save(temp.name)
        return temp.name


def soffice_available():
    """Check if soffice (LibreOffice) is available on the system."""
    return shutil.which("soffice") is not None


def test_inbound_with_docx(client):
    """Test that the inbound endpoint accepts .docx files using the real fixture."""
    fixture_path = Path("tests/fixtures/example-calendar.docx")
    assert fixture_path.exists(), "Fixture .docx file not found"
    with open(fixture_path, "rb") as f:
        data = {"attachment-1": (f, "example-calendar.docx")}
        response = client.post("/inbound", data=data)
        assert response.status_code == 200
        data = response.get_json()
        assert data["status"] == "success"
        assert "filename" in data
        assert len(data["events"]) > 0
        assert any("ADMIN DAY" in e["title"] for e in data["events"])


def test_inbound_with_doc(client):
    """Test that the inbound endpoint accepts .doc files using the real fixture, mocking soffice and creating the .docx file as a side effect."""
    fixture_path = Path("tests/fixtures/example-calendar.doc")
    assert fixture_path.exists(), "Fixture .doc file not found"
    # Save the .doc fixture to a temp file
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as temp_doc:
        with open(fixture_path, "rb") as f:
            temp_doc.write(f.read())
        temp_doc_path = temp_doc.name
    # Copy the .docx fixture to the expected path as a side effect of the mock
    docx_fixture = Path("tests/fixtures/example-calendar.docx")
    temp_docx_path = Path(temp_doc_path).with_suffix(".docx")
    with patch("subprocess.run") as mock_run:

        def side_effect(*args, **kwargs):
            shutil.copy(docx_fixture, temp_docx_path)
            return None

        mock_run.side_effect = side_effect
        # Debug logging
        print(f"Temp .doc path: {temp_doc_path}")
        print(f"Temp .docx path: {temp_docx_path}")
        print(f"Temp .doc exists: {os.path.exists(temp_doc_path)}")
        print(f"Temp .docx exists: {os.path.exists(temp_docx_path)}")
        try:
            with open(temp_doc_path, "rb") as f:
                data = {"attachment-1": (f, "example-calendar.doc")}
                response = client.post("/inbound", data=data)
                assert response.status_code == 200
                data = response.get_json()
                assert data["status"] == "success"
                assert "filename" in data
                assert len(data["events"]) > 0
                assert any("ADMIN DAY" in e["title"] for e in data["events"])
        finally:
            os.unlink(temp_doc_path)
            if temp_docx_path.exists():
                temp_docx_path.unlink()


def test_inbound_with_invalid_file(client):
    """Test that the inbound endpoint rejects invalid files. Mock soffice for .doc."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temp:
        temp.write(b"not a word document")
        temp_path = temp.name
    try:
        with open(temp_path, "rb") as f:
            data = {"attachment-1": (f, "test.docx")}
            response = client.post("/inbound", data=data)
            assert response.status_code == 400
            assert response.data.decode() == "Invalid .docx file"
        with patch("subprocess.run") as mock_run:
            mock_run.return_value = None  # Simulate successful conversion
            # Also, ensure a dummy .docx file exists after 'conversion'
            temp_dir = tempfile.gettempdir()
            dummy_docx = Path(temp_dir) / "test.docx"
            with open(dummy_docx, "wb") as f2:
                f2.write(b"not a word document")
            try:
                with open(temp_path, "rb") as f:
                    data = {"attachment-1": (f, "test.doc")}
                    response = client.post("/inbound", data=data)
                    assert response.status_code == 400
                    assert response.data.decode() == "Invalid .docx file"
            finally:
                if dummy_docx.exists():
                    dummy_docx.unlink()
    finally:
        os.unlink(temp_path)
